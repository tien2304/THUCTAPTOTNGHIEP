import openpyxl
from django.shortcuts import render, redirect, get_object_or_404
from django.core.paginator import Paginator
from django.contrib import messages
from django.db.models import Q
from django.contrib.auth.decorators import login_required
from django.urls import reverse
from Home.models import (
    KyThucTap, SinhVien, GiangVien, TaiLieu, NhiemVu, BaiNop, PhanCongGVHD, 
    HoiDong, HoiDong_GiangVien, HoiDong_SinhVien, BangDiem, TyLeDiem,
    ChiTietTraLoi, PhieuTraLoi
)
from Home.utils import sync_user_account, groups_required
from datetime import datetime

@groups_required('Giao vụ')
@login_required
def home_view(request):
    """View mặc định cho Giáo vụ – giờ đây sẽ chuyển hướng thẳng đến Kỳ thực tập."""
    return redirect('GiaoVu:giaovu_kythuctap')

# Quản lý kỳ thực tập
def ky_thuc_tap(request):
    if request.method == 'POST':
        ten_ky   = request.POST.get('ten', '').strip()
        bat_dau  = request.POST.get('bat_dau', '')
        ket_thuc = request.POST.get('ket_thuc', '')
        if ten_ky and bat_dau and ket_thuc:
            if bat_dau >= ket_thuc:
                messages.error(request, 'Ngày kết thúc phải lớn hơn ngày bắt đầu.')
                return redirect('GiaoVu:giaovu_kythuctap')

            KyThucTap.objects.create(
                ten_ky=ten_ky, ngay_bat_dau=bat_dau,
                ngay_ket_thuc=ket_thuc,
            )
            messages.success(request, f'Đã thêm kỳ thực tập "{ten_ky}" thành công!')
        else:
            messages.error(request, 'Vui lòng điền đầy đủ thông tin.')
        return redirect('GiaoVu:giaovu_kythuctap')

    tat_ca_ky = KyThucTap.objects.all().order_by('-id')
    so_trang = Paginator(tat_ca_ky, 10)
    trang_hien_tai = so_trang.get_page(request.GET.get('page', 1))
    context = {'current_page': 'kythuctap', 'trang_hien_tai': trang_hien_tai}
    return render(request, 'GiaoVu/ky_thuc_tap.html', context)

# Chỉnh sửa kỳ thực tập
def edit_ky_thuc_tap_view(request):
    if request.method == 'POST':
        ky_id = request.POST.get('id')
        ten_ky = request.POST.get('ten', '').strip()
        bat_dau = request.POST.get('bat_dau', '')
        ket_thuc = request.POST.get('ket_thuc', '')
        if ky_id and ten_ky and bat_dau and ket_thuc:
            # Kiểm tra tính hợp lệ của ngày
            if bat_dau >= ket_thuc:
                messages.error(request, 'Ngày kết thúc phải lớn hơn ngày bắt đầu.')
                return redirect('GiaoVu:giaovu_kythuctap')
            ky = KyThucTap.objects.filter(id=ky_id).first()
            if ky:
                ky.ten_ky = ten_ky
                ky.ngay_bat_dau = bat_dau
                ky.ngay_ket_thuc = ket_thuc
                ky.save()
                messages.success(request, f'Đã cập nhật kỳ thực tập "{ten_ky}" thành công!')
            else:
                messages.error(request, 'Không tìm thấy kỳ thực tập để cập nhật.')
        else:
            messages.error(request, 'Vui lòng điền đầy đủ thông tin.')
    return redirect('GiaoVu:giaovu_kythuctap')

# Quản lý tài liệu
def ql_tai_lieu_view(request):
    if request.method == 'POST':
        tieu_de = request.POST.get('tieu_de', '').strip()
        mo_ta = request.POST.get('mo_ta', '').strip()
        ky_id = request.POST.get('ky_id')
        file_dinh_kem = request.FILES.get('file_dinh_kem')

        if tieu_de and ky_id and file_dinh_kem:
            gv = GiangVien.objects.filter(ma_gv=request.user.username).first()
            ky = KyThucTap.objects.filter(id=ky_id).first()
            if gv and ky:
                TaiLieu.objects.create(
                    ten_tai_lieu=tieu_de,
                    mo_ta=mo_ta,
                    ky=ky,
                    duong_dan_file=file_dinh_kem,
                    gv_dang=gv
                )
                messages.success(request, f'Đã thêm tài liệu "{tieu_de}" thành công!')
            else:
                messages.error(request, 'Dữ liệu không hợp lệ (Giáo vụ sinh hoặc Kỳ thực tập không tồn tại).')
        else:
            messages.error(request, 'Vui lòng điền đầy đủ thông tin và chọn tệp đính kèm.')

        return redirect('GiaoVu:giaovu_tailieu')

    all_ky = KyThucTap.objects.all().order_by('-id')
    tai_lieu_list = TaiLieu.objects.all().order_by('-ngay_cap_nhat')

    paginator = Paginator(tai_lieu_list, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    context = {
        'current_page': 'tailieu',
        'page_obj': page_obj,
        'all_ky': all_ky,
    }
    return render(request, 'GiaoVu/tai_lieu.html', context)

# Chỉnh sửa tài liệu
def edit_tai_lieu_view(request):
    if request.method == 'POST':
        doc_id = request.POST.get('id')
        tieu_de = request.POST.get('tieu_de', '').strip()
        mo_ta = request.POST.get('mo_ta', '').strip()
        ky_id = request.POST.get('ky_id')
        file_dinh_kem = request.FILES.get('file_dinh_kem')

        if doc_id and tieu_de and ky_id:
            doc = TaiLieu.objects.filter(id=doc_id).first()
            ky = KyThucTap.objects.filter(id=ky_id).first()
            if doc and ky:
                doc.ten_tai_lieu = tieu_de
                doc.mo_ta = mo_ta
                doc.ky = ky
                if file_dinh_kem:
                    doc.duong_dan_file = file_dinh_kem
                doc.save()
                messages.success(request, f'Đã cập nhật tài liệu "{tieu_de}" thành công!')
            else:
                messages.error(request, 'Không tìm thấy tài liệu hoặc kỳ thực tập.')
        else:
            messages.error(request, 'Vui lòng điền đầy đủ thông tin bắt buộc.')

    return redirect('GiaoVu:giaovu_tailieu')

# Quản lý sinh viên
def ql_sinh_vien_view(request):
    if request.method == 'POST':
        ma_sv  = request.POST.get('ma_sv', '').strip()
        ho_ten = request.POST.get('ho_ten', '').strip()
        lop    = request.POST.get('lop', '').strip()
        ky_id  = request.POST.get('ky_hien_tai', '')
        if ma_sv and ho_ten and lop:
            ky = KyThucTap.objects.filter(id=ky_id).first() if ky_id else None
            sv = SinhVien.objects.create(ma_sv=ma_sv, ho_ten=ho_ten, lop=lop, ky_hien_tai=ky)
            
            mat_khau = request.POST.get('mat_khau', '').strip()
            if not mat_khau:
                mat_khau = "123456789"
                
            sync_user_account(ma_sv, ho_ten, 'SinhVien', password=mat_khau)
            messages.success(request, f'Đã thêm sinh viên "{ho_ten}" thành công!')
        else:
            messages.error(request, 'Vui lòng điền đầy đủ thông tin.')
        return redirect('GiaoVu:giaovu_qlsinhvien')

    selected_ky  = request.GET.get('ky', '')
    search_query = request.GET.get('q', '').strip()

    # Nếu chưa chọn kỳ, tự động chọn kỳ mới nhất
    if not selected_ky and 'ky' not in request.GET:
        newest = KyThucTap.objects.order_by('-id').first()
        if newest:
            selected_ky = str(newest.id)

    sv_list = SinhVien.objects.select_related('ky_hien_tai').all().order_by('ma_sv')
    if selected_ky:
        sv_list = sv_list.filter(ky_hien_tai__id=selected_ky)
    if search_query:
        sv_list = sv_list.filter(Q(ma_sv__icontains=search_query) | Q(ho_ten__icontains=search_query))
    # Đếm trang
    paginator = Paginator(sv_list, 10)
    page_obj  = paginator.get_page(request.GET.get('page', 1))
    context = {
        'current_page' : 'qlsinhvien',
        'danh_sach_ky' : KyThucTap.objects.all().order_by('-id'),
        'page_obj'     : page_obj,
        'selected_ky'  : selected_ky,
        'selected_ky_int': int(selected_ky) if selected_ky.isdigit() else None,
        'search_query' : search_query,
    }
    return render(request, 'GiaoVu/sinh_vien.html', context)

# Import sinh viên
def import_sinh_vien_view(request):
    """View xử lý import sinh viên từ file Excel."""
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']

        # 1. Lấy kỳ thực tập mới nhất
        ky_moi_nhat = KyThucTap.objects.order_by('-id').first()
        if not ky_moi_nhat:
            messages.error(request, 'Hệ thống chưa có Kỳ thực tập nào. Vui lòng tạo kỳ thực tập trước khi import.')
            return redirect('GiaoVu:giaovu_qlsinhvien')

        try:
            # 2. Đọc file Excel
            workbook = openpyxl.load_workbook(excel_file, data_only=True)
            sheet = workbook.active

            success_count = 0
            # Giả định: Cột A: MSSV, Cột B: Họ tên, Cột C: Lớp
            # Đọc từ hàng thứ 2 (bỏ qua tiêu đề)
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row[0]: continue  # Bỏ qua dòng trống MSSV

                ma_sv = str(row[0]).strip()

                ho_ten = str(row[1]).strip() if row[1] else "Chưa có tên"
                lop = str(row[2]).strip() if row[2] else "N/A"

                # 3. Lưu vào database (Update nếu trùng MSSV)
                SinhVien.objects.update_or_create(
                    ma_sv=ma_sv,
                    defaults={
                        'ho_ten': ho_ten,
                        'lop': lop,
                        'ky_hien_tai': ky_moi_nhat
                    }
                )
                sync_user_account(ma_sv, ho_ten, 'SinhVien')
                success_count += 1

            messages.success(request, f'Đã import thành công {success_count} sinh viên vào "{ky_moi_nhat.ten_ky}".')
        except Exception as e:
            messages.error(request, f'Lỗi khi xử lý file: {str(e)}')

    return redirect('GiaoVu:giaovu_qlsinhvien')


def edit_sinh_vien_view(request):
    """View xử lý cập nhật thông tin sinh viên (bao gồm cả việc đổi MSSV)."""
    if request.method == 'POST':
        original_ma = request.POST.get('original_ma_sv', '').strip()
        new_ma = request.POST.get('ma_sv', '').strip()
        ho_ten = request.POST.get('ho_ten', '').strip()
        lop = request.POST.get('lop', '').strip()
        ky_id = request.POST.get('ky_hien_tai', '')

        if original_ma and new_ma and ho_ten and lop:
            sinh_vien = SinhVien.objects.filter(ma_sv=original_ma).first()
            if sinh_vien:
                ky = KyThucTap.objects.filter(id=ky_id).first() if ky_id else None
                if original_ma != new_ma:
                    # Nếu đổi MSSV (Primary Key)
                    # 1. Tạo bản ghi mới dựa trên bản ghi cũ
                    SinhVien.objects.create(
                        ma_sv=new_ma,
                        ho_ten=ho_ten,
                        lop=lop,
                        ky_hien_tai=ky
                    )
                    # 2. Xóa bản ghi cũ
                    sinh_vien.delete()
                else:
                    # Nếu giữ nguyên MSSV
                    sinh_vien.ho_ten = ho_ten
                    sinh_vien.lop = lop
                    sinh_vien.ky_hien_tai = ky
                    sinh_vien.save()

                # Cập nhật tài khoản (và mật khẩu nếu có)
                mat_khau_moi = request.POST.get('mat_khau', '').strip()
                if mat_khau_moi:
                    sync_user_account(new_ma, ho_ten, 'SinhVien', password=mat_khau_moi)
                else:
                    # Chỉ cập nhật họ tên nếu không đổi mật khẩu
                    sync_user_account(new_ma, ho_ten, 'SinhVien')

                messages.success(request, f'Cập nhật sinh viên "{ho_ten}" thành công!')
            else:
                messages.error(request, 'Không tìm thấy sinh viên để cập nhật.')
        else:
            messages.error(request, 'Vui lòng điền đầy đủ thông tin.')

    return redirect('GiaoVu:giaovu_qlsinhvien')

def delete_sinh_vien_view(request, ma_sv):
    sinh_vien = SinhVien.objects.filter(ma_sv=ma_sv).first()
    if sinh_vien:
        ho_ten = sinh_vien.ho_ten
        sinh_vien.delete()
        messages.success(request, f'Đã xóa sinh viên "{ho_ten}" thành công!')
    else:
        messages.error(request, 'Không tìm thấy sinh viên để xóa.')
    return redirect('GiaoVu:giaovu_qlsinhvien')

def ql_giang_vien_view(request):
    """Trang Quản lý giảng viên – Giáo Vụ."""
    if request.method == 'POST':
        ma_gv       = request.POST.get('ma_gv', '').strip()
        ho_ten      = request.POST.get('ho_ten', '').strip()
        hoc_vi      = request.POST.get('hoc_vi', '').strip()
        chuc_vu     = request.POST.get('chuc_vu', '').strip()
        # Lấy danh sách chuyên môn từ checkboxes
        chuyen_mon_list = request.POST.getlist('chuyen_mon')
        chuyen_mon      = ", ".join(chuyen_mon_list) # Nối thành chuỗi để lưu
        sdt         = request.POST.get('so_dien_thoai', '').strip()

        if ma_gv and ho_ten and chuyen_mon:
            GiangVien.objects.create(
                ma_gv=ma_gv, ho_ten=ho_ten, hoc_vi=hoc_vi,
                chuc_vu=chuc_vu, chuyen_mon=chuyen_mon, so_dien_thoai=sdt,
            )
            # Map chuc_vu sang VaiTro
            role_map = {'Giảng viên': 'Giảng viên', 'Giáo vụ': 'Giáo vụ', 'Trưởng bộ môn': 'Trưởng bộ môn'}
            role_name = role_map.get(chuc_vu, 'Giảng viên')
            
            mat_khau = request.POST.get('mat_khau', '').strip()
            if not mat_khau:
                mat_khau = "123456789"
                
            sync_user_account(ma_gv, ho_ten, role_name, password=mat_khau)
            
            messages.success(request, f'Đã thêm giảng viên "{ho_ten}" thành công!')
        else:
            messages.error(request, 'Vui lòng điền đầy đủ thông tin bắt buộc.')
        return redirect('GiaoVu:giaovu_qlgiangvien')

    search_query = request.GET.get('q', '').strip()
    gv_list = GiangVien.objects.all().order_by('ho_ten')
    if search_query:
        gv_list = gv_list.filter(
            Q(ho_ten__icontains=search_query) | Q(chuyen_mon__icontains=search_query)
        )

    paginator = Paginator(gv_list, 10)
    page_obj  = paginator.get_page(request.GET.get('page', 1))

    # Lấy các danh mục từ Model để hiển thị lên Form
    hoc_vi_choices = [c[0] for c in GiangVien.HOC_VI_CHOICES]
    chuc_vu_choices = [c[0] for c in GiangVien.CHUC_VU_CHOICES]
    chuyen_mon_choices = [c[0] for c in GiangVien.CHUYEN_MON_CHOICES]

    context = {
        'current_page' : 'qlgiangvien',
        'page_obj'     : page_obj,
        'search_query' : search_query,
        'hoc_vi_choices': hoc_vi_choices,
        'chuc_vu_choices': chuc_vu_choices,
        'chuyen_mon_choices': chuyen_mon_choices,
    }
    return render(request, 'GiaoVu/giang_vien.html', context)


def import_giang_vien_view(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        try:
            workbook = openpyxl.load_workbook(excel_file, data_only=True)
            sheet = workbook.active
            success_count = 0
            for row in sheet.iter_rows(min_row=2, values_only=True):
                ma_gv_raw = row[0]
                if not ma_gv_raw: continue
                # Làm sạch MaGV
                ma_gv = str(ma_gv_raw).strip()
                ho_ten      = str(row[1]).strip() if row[1] else "Chưa rõ"
                hoc_vi      = str(row[2]).strip() if row[2] else "Thạc sĩ"
                chuc_vu     = str(row[3]).strip() if row[3] else "Giảng viên"
                chuyen_mon  = str(row[4]).strip() if row[4] else ""
                sdt = str(row[5]).strip() if row[5] else ""

                GiangVien.objects.update_or_create(
                    ma_gv=ma_gv,
                    defaults={
                        'ho_ten': ho_ten,
                        'hoc_vi': hoc_vi,
                        'chuc_vu': chuc_vu,
                        'chuyen_mon': chuyen_mon,
                        'so_dien_thoai': sdt
                    }
                )
                # Map chuc_vu sang VaiTro
                role_map = {'Giảng viên': 'Giảng viên', 'Giáo vụ': 'Giáo vụ', 'Trưởng bộ môn': 'Trưởng bộ môn'}
                role_name = role_map.get(chuc_vu, 'Giảng viên')
                sync_user_account(ma_gv, ho_ten, role_name)
                
                success_count += 1

            messages.success(request, f'Đã import thành công {success_count} giảng viên.')
        except Exception as e:
            messages.error(request, f'Lỗi khi xử lý file: {str(e)}')

    return redirect('GiaoVu:giaovu_qlgiangvien')


def edit_giang_vien_view(request):
    """View xử lý cập nhật giảng viên (Hỗ trợ đổi MaGV)."""
    if request.method == 'POST':
        original_ma = request.POST.get('original_ma_gv', '').strip()
        new_ma      = request.POST.get('ma_gv', '').strip()
        ho_ten      = request.POST.get('ho_ten', '').strip()
        hoc_vi      = request.POST.get('hoc_vi', '').strip()
        chuc_vu     = request.POST.get('chuc_vu', '').strip()
        sdt         = request.POST.get('so_dien_thoai', '').strip()
        
        # Chuyên môn từ Multi-select checkbox
        chuyen_mon_list = request.POST.getlist('chuyen_mon')
        chuyen_mon      = ", ".join(chuyen_mon_list)

        if original_ma and new_ma and ho_ten:
            gv = GiangVien.objects.filter(ma_gv=original_ma).first()
            if gv:
                if original_ma != new_ma:
                    # Đổi mã giảng viên (PK)
                    GiangVien.objects.create(
                        ma_gv=new_ma, ho_ten=ho_ten, hoc_vi=hoc_vi,
                        chuc_vu=chuc_vu, chuyen_mon=chuyen_mon, so_dien_thoai=sdt
                    )
                    gv.delete()
                else:
                    gv.ho_ten = ho_ten
                    gv.hoc_vi = hoc_vi
                    gv.chuc_vu = chuc_vu
                    gv.chuyen_mon = chuyen_mon
                    gv.so_dien_thoai = sdt
                    gv.save()
                
                # Cập nhật tài khoản (và mật khẩu nếu có)
                role_map = {'Giảng viên': 'Giảng viên', 'Giáo vụ': 'Giáo vụ', 'Trưởng bộ môn': 'Trưởng bộ môn'}
                role_name = role_map.get(chuc_vu, 'Giảng viên')
                
                mat_khau_moi = request.POST.get('mat_khau', '').strip()
                if mat_khau_moi:
                    sync_user_account(new_ma, ho_ten, role_name, password=mat_khau_moi)
                else:
                    sync_user_account(new_ma, ho_ten, role_name)

                messages.success(request, f'Cập nhật giảng viên "{ho_ten}" thành công!')
            else:
                messages.error(request, 'Không tìm thấy giảng viên.')
        else:
            messages.error(request, 'Thiếu thông tin bắt buộc.')
    return redirect('GiaoVu:giaovu_qlgiangvien')


def delete_giang_vien_view(request, ma_gv):
    """Xóa giảng viên."""
    gv = GiangVien.objects.filter(ma_gv=ma_gv).first()
    if gv:
        ho_ten = gv.ho_ten
        gv.delete()
        messages.success(request, f'Đã xóa giảng viên "{ho_ten}".')
    else:
        messages.error(request, 'Không tìm thấy giảng viên.')
    return redirect('GiaoVu:giaovu_qlgiangvien')

def nhiem_vu_view(request):
    if request.method == 'POST':
        ky_id = request.POST.get('ky_id')
        ten_nhiem_vu = request.POST.get('ten_nhiem_vu', '').strip()
        mo_ta = request.POST.get('mo_ta', '').strip()
        han_nop = request.POST.get('han_nop')
        
        if ky_id and ten_nhiem_vu and han_nop:
            ky = KyThucTap.objects.filter(id=ky_id).first()
            if ky:
                try:
                    han_nop_dt = datetime.strptime(han_nop, '%Y-%m-%dT%H:%M')
                    NhiemVu.objects.create(
                        ky=ky,
                        ten_nhiem_vu=ten_nhiem_vu,
                        han_nop=han_nop_dt,
                        mo_ta=mo_ta
                    )
                    messages.success(request, f'Đã thiết lập nhiệm vụ "{ten_nhiem_vu}" thành công!')
                except Exception as e:
                    messages.error(request, f'Lỗi định dạng ngày giờ: {str(e)}')
            else:
                messages.error(request, 'Kỳ thực tập không tồn tại.')
        else:
            messages.error(request, 'Vui lòng điền đủ thông tin tên nhiệm vụ, mô tả và hạn nộp.')
            
        return redirect(f"/giao-vu/nhiem-vu/?ky_id={ky_id}" if ky_id else "GiaoVu:giaovu_nhiemvu")

    all_ky = KyThucTap.objects.all().order_by('-id')
    ky_id = request.GET.get('ky_id')
    
    if ky_id:
        selected_ky = KyThucTap.objects.filter(id=ky_id).first()
    else:
        selected_ky = all_ky.first()
        
    nhiem_vu_list = []
    if selected_ky:
        nhiem_vu_list = NhiemVu.objects.filter(ky=selected_ky).order_by('han_nop')
        
    context = {
        'current_page': 'nhiemvu',
        'all_ky': all_ky,
        'selected_ky': selected_ky,
        'nhiem_vu_list': nhiem_vu_list,
    }
    return render(request, 'GiaoVu/nhiem_vu.html', context)

def giang_vien_hd_view(request):
    """Trang xem Giảng viên hướng dẫn – Giáo Vụ."""
    all_ky = KyThucTap.objects.all().order_by('-id')
    ky_id = request.GET.get('ky_id')
    
    if ky_id:
        selected_ky = KyThucTap.objects.filter(id=ky_id).first()
    else:
        selected_ky = all_ky.first()
        
    ds_sinh_vien = []
    if selected_ky:
        # Lấy toàn bộ phân công ĐÃ DUYỆT của kỳ này
        phan_cong_qs = PhanCongGVHD.objects.filter(ky=selected_ky, trang_thai=2).select_related('giang_vien', 'sinh_vien').order_by('sinh_vien__ma_sv')
        
        # Build map prefix học vị
        PREFIX_MAP = {'Thạc sĩ': 'ThS.', 'Tiến sĩ': 'TS.', 'Phó Giáo sư': 'PGS.TS.', 'Giáo sư': 'GS.TS.'}
        
        for pc in phan_cong_qs:
            gv = pc.giang_vien
            pf = PREFIX_MAP.get(gv.hoc_vi, '')
            ten_gv_full = f"{pf} {gv.ho_ten}".strip() if pf else gv.ho_ten
            
            ds_sinh_vien.append({
                'ma_sv': pc.sinh_vien.ma_sv,
                'ho_ten': pc.sinh_vien.ho_ten,
                'lop': pc.sinh_vien.lop,
                'ten_gvhd': ten_gv_full,
                'trang_thai': pc.trang_thai
            })

    context = {
        'current_page': 'giangvienhd',
        'all_ky': all_ky,
        'selected_ky': selected_ky,
        'selected_ky_id': selected_ky.id if selected_ky else None,
        'ds_sinh_vien': ds_sinh_vien,
    }
    return render(request, 'GiaoVu/giang_vien_hd.html', context)

def hoidong_view(request):
    """Trang Quản lý Hội đồng – Giáo Vụ."""
    all_ky = KyThucTap.objects.all().order_by('-id')
    ky_id = request.GET.get('ky_id')
    
    if ky_id:
        selected_ky = KyThucTap.objects.filter(id=ky_id).first()
    else:
        selected_ky = all_ky.first()
        
    ds_hoidong = []
    if selected_ky:
        # Lấy danh sách hội đồng của kỳ này (Chỉ lấy những cái ĐÃ DUYỆT)
        hoidong_qs = HoiDong.objects.filter(ky=selected_ky, trang_thai=2).order_by('ngay_bao_ve', 'thoi_gian_bat_dau')
        
        for hd in hoidong_qs:
            # Đếm số giảng viên và sinh viên trong hội đồng này
            gv_list = HoiDong_GiangVien.objects.filter(hoi_dong=hd).select_related('giang_vien')
            sv_list = HoiDong_SinhVien.objects.filter(hoi_dong=hd).select_related('sinh_vien')
            
            # Kiểm tra xem đã đủ thông tin chưa
            is_ready = hd.thoi_gian_bat_dau and hd.thoi_gian_ket_thuc and hd.dia_diem
            
            ds_hoidong.append({
                'id': hd.id,
                'ten': hd.ten_hoi_dong,
                'thoi_gian_bat_dau': hd.thoi_gian_bat_dau,
                'thoi_gian_ket_thuc': hd.thoi_gian_ket_thuc,
                'thoi_gian': f"{hd.thoi_gian_bat_dau.strftime('%H:%M')} - {hd.thoi_gian_ket_thuc.strftime('%H:%M')}" if is_ready else "Chưa thiết lập",
                'ngay_bao_ve': hd.ngay_bao_ve,
                'dia_diem': hd.dia_diem,
                'gv_count': gv_list.count(),
                'sv_count': sv_list.count(),
                'giang_vien': [g.giang_vien for g in gv_list],
                'sinh_vien': [s.sinh_vien for s in sv_list],
                'trang_thai': hd.trang_thai
            })
            
    context = {
        'current_page': 'hoidong',
        'all_ky': all_ky,
        'selected_ky': selected_ky,
        'selected_ky_id': selected_ky.id if selected_ky else None,
        'ds_hoidong': ds_hoidong,
    }
    return render(request, 'GiaoVu/hoi_dong.html', context)

def chi_tiet_hoidong_view(request, hd_id):
    """Trang chi tiết của một Hội đồng – Giáo Vụ."""
    hoidong = HoiDong.objects.filter(id=hd_id).select_related('ky').first()
    
    if not hoidong:
        messages.error(request, "Không tìm thấy hội đồng này.")
        return redirect('GiaoVu:giaovu_hoidong')
        
    ds_giang_vien = HoiDong_GiangVien.objects.filter(hoi_dong=hoidong).select_related('giang_vien')
    ds_sinh_vien_raw = HoiDong_SinhVien.objects.filter(hoi_dong=hoidong).select_related('sinh_vien')
    
    # Lấy thông tin GVHD cho từng sinh viên
    ds_sinh_vien = []
    for item in ds_sinh_vien_raw:
        sv = item.sinh_vien
        # Tìm GVHD trong kỳ này
        pc = PhanCongGVHD.objects.filter(sinh_vien=sv, ky=hoidong.ky, trang_thai=2).select_related('giang_vien').first()
        ds_sinh_vien.append({
            'sinh_vien': sv,
            'gvhd': pc.giang_vien if pc else None
        })
    
    context = {
        'current_page': 'hoidong',
        'hoidong': hoidong,
        'ds_giang_vien': ds_giang_vien,
        'ds_sinh_vien': ds_sinh_vien,
    }
    return render(request, 'GiaoVu/chi_tiet_hoidong.html', context)
from datetime import time
from django.http import JsonResponse
import json
from django.contrib.auth.decorators import login_required
from django.shortcuts import get_object_or_404

@login_required
def cap_nhat_thoi_gian_hoi_dong(request, hd_id):
    hoidong = get_object_or_404(HoiDong, id=hd_id)

    if hoidong.trang_thai != 2:
        return JsonResponse({"status": "error", "message": "Hội đồng chưa được duyệt."}, status=400)

    if request.method != "POST":
        return JsonResponse({"status": "error", "message": "Invalid method"}, status=405)

    try:
        data = json.loads(request.body)
        time_range = data.get("time_range", "").strip()
        dia_diem = data.get("dia_diem", "").strip()
        if time_range and time_range != "":
            bat_dau = None
            ket_thuc = None
            try:
                if '-' in time_range:
                    start_str, end_str = [x.strip() for x in time_range.split('-', 1)]
                    h1, m1 = map(int, start_str.split(':'))
                    h2, m2 = map(int, end_str.split(':'))
                    bat_dau = time(hour=h1, minute=m1)
                    ket_thuc = time(hour=h2, minute=m2)
                else:
                    h, m = map(int, time_range.split(':'))
                    bat_dau = time(hour=h, minute=m)
                
                hoidong.thoi_gian_bat_dau = bat_dau
                hoidong.thoi_gian_ket_thuc = ket_thuc
            except Exception as e:
                return JsonResponse({"status": "error", "message": f"Lỗi định dạng thời gian: {str(e)}"}, status=400)

        hoidong.dia_diem = dia_diem
        hoidong.save()

        return JsonResponse({
            "status": "success",
            "message": "Cập nhật thành công!"
        })

    except Exception as e:
        return JsonResponse({"status": "error", "message": str(e)}, status=500)

from django.contrib.auth.decorators import login_required
from django.contrib.auth import update_session_auth_hash

@login_required
def update_password(request):
    if request.method == "POST":
        current_password = request.POST.get('current_password')
        new_password = request.POST.get('new_password')
        confirm_password = request.POST.get('confirm_password')

        if not request.user.check_password(current_password):
            messages.error(request, "Mật khẩu hiện tại không đúng.", extra_tags='pwd_error')
            return redirect('GiaoVu:giaovu_home')
        
        if new_password != confirm_password:
            messages.error(request, "Xác nhận mật khẩu không khớp.", extra_tags='pwd_error')
            return redirect('GiaoVu:giaovu_home')
        
        request.user.set_password(new_password)
        request.user.save()
        update_session_auth_hash(request, request.user)
        messages.success(request, "Thay đổi mật khẩu thành công!", extra_tags='pwd_success')
        
    return redirect('GiaoVu:giaovu_home')

def ql_diem_view(request):
    """Trang Quản lý điểm của toàn bộ sinh viên – Giáo Vụ."""
    selected_ky = request.GET.get('ky', '')
    
    # Nếu chưa chọn kỳ, tự động chọn kỳ mới nhất
    if not selected_ky and 'ky' not in request.GET:
        newest = KyThucTap.objects.order_by('-id').first()
        if newest:
            selected_ky = str(newest.id)

    # Lấy danh sách điểm
    diem_list = BangDiem.objects.select_related('sinh_vien', 'ky').all().order_by('sinh_vien__ma_sv')
    
    if selected_ky:
        diem_list = diem_list.filter(ky__id=selected_ky)

    # Lấy thêm thông tin doanh nghiệp từ khảo sát
    for d in diem_list:
        # Tìm các câu trả lời liên quan đến doanh nghiệp của SV này trong kỳ này
        answers = ChiTietTraLoi.objects.filter(
            phieu_tra_loi__sinh_vien=d.sinh_vien,
            phieu_tra_loi__mau_khao_sat__ky=d.ky
        ).select_related('cau_hoi')
        
        d.enterprise_info = {
            'ten': (d.sinh_vien.noi_thuc_tap or "").strip(),
            'sdt': "",
            'dia_chi': "",
            'email': ""
        }
        
        for ans in answers:
            tag = (ans.cau_hoi.system_tag or "").strip().lower()
            val = (ans.gia_tri or "").strip()
            if not val: continue

            if tag == "don_vi_tt" and not d.enterprise_info['ten']:
                d.enterprise_info['ten'] = val
            elif tag == "sdt_don_vi":
                d.enterprise_info['sdt'] = val
            elif tag == "dia_diem_dv":
                d.enterprise_info['dia_chi'] = val
            elif tag == "email_don_vi":
                d.enterprise_info['email'] = val
            elif tag in ["diem_doanh_nghiep", "diem_dn"]:
                try:
                    if d.diem_doanh_nghiep is None:
                        d.diem_doanh_nghiep = float(val)
                except (ValueError, TypeError):
                    pass



    paginator = Paginator(diem_list, 15)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    
    context = {
        'current_page': 'diem',
        'danh_sach_ky': KyThucTap.objects.all().order_by('-id'),
        'page_obj': page_obj,
        'selected_ky': selected_ky,
        'selected_ky_int': int(selected_ky) if selected_ky.isdigit() else None,
    }
    return render(request, 'GiaoVu/diem.html', context)


def cap_nhat_diem_dn(request, ma_sv):
    """Giáo vụ cập nhật điểm doanh nghiệp."""
    if request.method == 'POST':
        diem_dn = request.POST.get('diem_dn', '').strip()
        ky_id = request.POST.get('ky_id')
        
        if diem_dn:
            try:
                # Xử lý dấu phẩy sang dấu chấm nếu người dùng nhập kiểu VN
                diem_dn = diem_dn.replace(',', '.')
                diem_dn_val = float(diem_dn)
                
                if not (0 <= diem_dn_val <= 10):
                    messages.error(request, "Điểm phải nằm trong khoảng từ 0 đến 10.")
                else:
                    sv = SinhVien.objects.get(ma_sv=ma_sv)
                    ky = KyThucTap.objects.get(id=ky_id)
                    bd, _ = BangDiem.objects.get_or_create(sinh_vien=sv, ky=ky)
                    bd.diem_doanh_nghiep = diem_dn_val
                    bd.save()
                    if hasattr(bd, 'calculate_total'):
                        bd.calculate_total()
            except Exception as e:
                messages.error(request, f"Lỗi cập nhật: {str(e)}")
        else:
             messages.error(request, "Vui lòng nhập điểm.")

    return redirect(f"{reverse('GiaoVu:giaovu_diem')}?ky={ky_id}")

def edit_nhiem_vu_view(request):
    """View xử lý cập nhật nhiệm vụ."""
    if request.method == 'POST':
        nv_id = request.POST.get('id')
        ten_nhiem_vu = request.POST.get('ten_nhiem_vu', '').strip()
        mo_ta = request.POST.get('mo_ta', '').strip()
        ky_id = request.POST.get('ky_id')
        han_nop = request.POST.get('han_nop')

        if nv_id and ten_nhiem_vu and ky_id and han_nop:
            nv = NhiemVu.objects.filter(id=nv_id).first()
            ky = KyThucTap.objects.filter(id=ky_id).first()
            if nv and ky:
                try:
                    han_nop_dt = datetime.strptime(han_nop, '%Y-%m-%dT%H:%M')
                    nv.ten_nhiem_vu = ten_nhiem_vu
                    nv.mo_ta = mo_ta
                    nv.ky = ky
                    nv.han_nop = han_nop_dt
                    nv.save()
                    messages.success(request, f'Đã cập nhật nhiệm vụ "{ten_nhiem_vu}" thành công!')
                except Exception as e:
                    messages.error(request, f'Lỗi định dạng ngày giờ: {str(e)}')
            else:
                messages.error(request, 'Không tìm thấy nhiệm vụ hoặc kỳ thực tập.')
        else:
            messages.error(request, 'Vui lòng điền đầy đủ thông tin bắt buộc.')
            
    return redirect(f"/giao-vu/nhiem-vu/?ky_id={ky_id}" if ky_id else "GiaoVu:giaovu_nhiemvu")

def delete_nhiem_vu_view(request, nv_id):
    """Xóa nhiệm vụ."""
    nv = NhiemVu.objects.filter(id=nv_id).first()
    ky_id = None
    if nv:
        ky_id = nv.ky.id
        ten = nv.ten_nhiem_vu
        nv.delete()
        messages.success(request, f'Đã xóa nhiệm vụ "{ten}" thành công!')
    else:
        messages.error(request, 'Không tìm thấy nhiệm vụ để xóa.')
        
    return redirect(f"/giao-vu/nhiem-vu/?ky_id={ky_id}" if ky_id else "GiaoVu:giaovu_nhiemvu")

def chi_tiet_nhiem_vu_view(request, nv_id):
    """Trang chi tiết nhiệm vụ - xem danh sách nộp bài của sinh viên."""
    nhiem_vu = get_object_or_404(NhiemVu, id=nv_id)
    ky = nhiem_vu.ky
    
    # Lấy danh sách toàn bộ sinh viên trong kỳ này
    sinh_vien_list = SinhVien.objects.filter(ky_hien_tai=ky).order_by('ma_sv')
    
    # Lấy danh sách bài nộp của nhiệm vụ này
    bai_nop_list = BaiNop.objects.filter(nhiem_vu=nhiem_vu).select_related('sinh_vien')
    bai_nop_dict = {bn.sinh_vien.ma_sv: bn for bn in bai_nop_list}
    
    # Kết hợp dữ liệu
    data = []
    for sv in sinh_vien_list:
        bn = bai_nop_dict.get(sv.ma_sv)
        data.append({
            'sinh_vien': sv,
            'bai_nop': bn,
            'da_nop': True if bn else False
        })
        
    # Xử lý lọc
    filter_status = request.GET.get('status', 'all')
    if filter_status == 'submitted':
        data = [d for d in data if d['da_nop']]
    elif filter_status == 'pending':
        data = [d for d in data if not d['da_nop']]
        
    context = {
        'current_page': 'nhiemvu',
        'nhiem_vu': nhiem_vu,
        'data': data,
        'filter_status': filter_status,
        'total_count': len(data),
        'submitted_count': len([d for d in data if d['da_nop']]),
    }
    return render(request, 'GiaoVu/chi_tiet_nhiem_vu.html', context)


